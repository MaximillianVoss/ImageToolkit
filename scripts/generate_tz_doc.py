from __future__ import annotations

from copy import deepcopy
from datetime import date
from pathlib import Path

import win32com.client
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = Path(
    r"c:\Users\FossW\OneDrive\Документы\Настраиваемые шаблоны Office\Шаблон пояснительной записки с макросами.dotm"
)
TMP_DIR = ROOT / "tmp" / "docs"
OUTPUT_DIR = ROOT / "output" / "doc"
BASE_DOCX_PATH = TMP_DIR / "template_base.docx"
OUTPUT_DOCX_PATH = OUTPUT_DIR / (
    "\u0422\u0435\u0445\u043d\u0438\u0447\u0435\u0441\u043a\u043e\u0435 \u0437\u0430\u0434\u0430\u043d\u0438\u0435 ImageToolkit.docx"
)
OUTPUT_PDF_PATH = OUTPUT_DIR / (
    "\u0422\u0435\u0445\u043d\u0438\u0447\u0435\u0441\u043a\u043e\u0435 \u0437\u0430\u0434\u0430\u043d\u0438\u0435 ImageToolkit.pdf"
)

TEXT_STYLE = "ГОСТ Текст"
H1_STYLE = "ГОСТ Заголовок 1"
H2_STYLE = "ГОСТ Заголовок 2"
TABLE_STYLE = "ГОСТ.Черный"


def ensure_directories() -> None:
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def create_base_document_from_template() -> None:
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0

    try:
        document = word.Documents.Add(str(TEMPLATE_PATH))
        document.SaveAs2(str(BASE_DOCX_PATH), FileFormat=12)
        document.Close(False)
    finally:
        word.Quit()


def clear_document_body(document: Document) -> None:
    body = document._element.body
    sect_pr = body.sectPr

    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_default_font_size(paragraph, size_pt: int) -> None:
    for run in paragraph.runs:
        run.font.size = Pt(size_pt)


def add_text_paragraph(document: Document, text: str = ""):
    paragraph = document.add_paragraph(style=TEXT_STYLE)
    if text:
        paragraph.add_run(text)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    style = H1_STYLE if level == 1 else H2_STYLE
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        add_text_paragraph(document, f"- {item}")


def add_toc_field(document: Document) -> None:
    paragraph = document.add_paragraph(style=TEXT_STYLE)
    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")

    hint = OxmlElement("w:t")
    hint.text = "Содержание будет обновлено при открытии документа."

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(hint)
    run._r.append(fld_end)


def add_page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break()


def clone_last_section(document: Document) -> None:
    new_section = document.add_section(WD_SECTION.NEW_PAGE)
    source_section = document.sections[0]
    target_section = document.sections[-1]
    target_section.top_margin = source_section.top_margin
    target_section.bottom_margin = source_section.bottom_margin
    target_section.left_margin = source_section.left_margin
    target_section.right_margin = source_section.right_margin
    target_section.page_width = source_section.page_width
    target_section.page_height = source_section.page_height
    target_section.different_first_page_header_footer = source_section.different_first_page_header_footer
    if new_section.start_type != WD_SECTION.NEW_PAGE:
        pass


def add_title_page(document: Document) -> None:
    paragraph = add_text_paragraph(document)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("ТЕХНИЧЕСКОЕ ЗАДАНИЕ").bold = True
    set_default_font_size(paragraph, 16)

    paragraph = add_text_paragraph(document)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("на разработку программного средства для обработки изображений").bold = True
    set_default_font_size(paragraph, 14)

    paragraph = add_text_paragraph(document)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("ImageToolkit").bold = True
    set_default_font_size(paragraph, 14)

    for _ in range(6):
        add_text_paragraph(document)

    meta = [
        "Объект разработки: настольное приложение для обработки изображений.",
        "Основание: учебное задание по теме «создать программное средство для обработки изображений».",
        f"Дата составления: {date.today().strftime('%d.%m.%Y')}.",
    ]
    for item in meta:
        paragraph = add_text_paragraph(document, item)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for _ in range(10):
        add_text_paragraph(document)

    paragraph = add_text_paragraph(document, "Москва 2026")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_sources_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = TABLE_STYLE
    header = table.rows[0].cells
    header[0].text = "Обозначение"
    header[1].text = "Наименование"
    header[2].text = "Примечание"

    rows = [
        (
            "ГОСТ 19.201-78",
            "Техническое задание. Требования к содержанию и оформлению",
            "Основной норматив для структуры документа",
        ),
        (
            "ГОСТ 19.101-77",
            "Виды программ и программных документов",
            "Определяет состав программной документации",
        ),
        (
            "Задание на проект",
            "Создать программное средство для обработки изображений",
            "Исходное учебное требование",
        ),
    ]

    for designation, name, note in rows:
        cells = table.add_row().cells
        cells[0].text = designation
        cells[1].text = name
        cells[2].text = note


def add_stages_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = TABLE_STYLE
    header = table.rows[0].cells
    header[0].text = "Этап"
    header[1].text = "Содержание работ"
    header[2].text = "Результат"

    rows = [
        (
            "1. Аналитический",
            "Сбор требований, определение состава операций обработки изображений, выбор технологического стека.",
            "Уточненные требования и архитектурные решения.",
        ),
        (
            "2. Проектный",
            "Проектирование интерфейса, модели данных изображения и набора преобразований.",
            "Проект приложения и описание пользовательских сценариев.",
        ),
        (
            "3. Реализация",
            "Разработка WPF-интерфейса, библиотеки обработки изображений и тестового проекта.",
            "Работоспособный программный продукт и исходный код.",
        ),
        (
            "4. Тестирование",
            "Проверка сборки, модульных тестов и основных пользовательских сценариев.",
            "Подтверждение корректности ключевых функций.",
        ),
        (
            "5. Документирование",
            "Подготовка технического задания, пояснительной записки и сопроводительных материалов.",
            "Комплект программной документации.",
        ),
    ]

    for stage, content, result in rows:
        cells = table.add_row().cells
        cells[0].text = stage
        cells[1].text = content
        cells[2].text = result


def add_main_content(document: Document) -> None:
    paragraph = document.add_paragraph(style=TEXT_STYLE)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Содержание")
    run.bold = True
    run.font.size = Pt(14)
    add_toc_field(document)
    clone_last_section(document)

    add_heading(document, "Введение", level=1)
    add_text_paragraph(
        document,
        "Настоящее техническое задание определяет состав, назначение, требования к функциональным и эксплуатационным характеристикам, "
        "а также порядок разработки и приемки программного средства для обработки изображений ImageToolkit."
    )
    add_text_paragraph(
        document,
        "Документ предназначен для использования в качестве базового регламентирующего материала при разработке, тестировании и "
        "оформлении сопроводительной документации по проекту."
    )

    add_heading(document, "Основания для разработки", level=1)
    add_text_paragraph(
        document,
        "Основанием для разработки является учебное задание на тему «создать программное средство для обработки изображений»."
    )
    add_text_paragraph(
        document,
        "Разработка ведется как самостоятельный программный проект, включающий настольное приложение, библиотеку алгоритмов обработки "
        "изображений и отдельный проект модульных тестов."
    )
    add_heading(document, "Источники и нормативные материалы", level=2)
    add_sources_table(document)

    add_heading(document, "Назначение разработки", level=1)
    add_heading(document, "Функциональное назначение", level=2)
    add_text_paragraph(
        document,
        "Программное средство предназначено для загрузки растровых изображений, применения к ним типовых операций обработки и сохранения "
        "полученного результата в графический файл."
    )
    add_heading(document, "Эксплуатационное назначение", level=2)
    add_text_paragraph(
        document,
        "Программа предназначена для использования обучающимися, преподавателями и иными пользователями, которым требуется простое "
        "настольное приложение для демонстрации и выполнения базовой обработки изображений в среде Microsoft Windows."
    )

    add_heading(document, "Требования к программе", level=1)
    add_heading(document, "Требования к функциональным характеристикам", level=2)
    add_text_paragraph(document, "Программа должна обеспечивать выполнение следующих функций:")
    add_bullets(
        document,
        [
            "загрузка изображения из файла локальной файловой системы;",
            "отображение исходного изображения и результата обработки в пользовательском интерфейсе;",
            "выбор операции обработки из предопределенного списка;",
            "предпросмотр результата до окончательного применения операции;",
            "применение операции к рабочей копии изображения;",
            "сброс рабочей копии к исходному состоянию;",
            "сохранение обработанного изображения в файл;",
            "отображение служебной информации о размерах изображения и текущем состоянии обработки.",
        ],
    )
    add_text_paragraph(document, "Минимальный набор операций обработки должен включать:")
    add_bullets(
        document,
        [
            "перевод изображения в оттенки серого;",
            "инверсию цветов;",
            "изменение яркости с регулируемым параметром;",
            "пороговую обработку с регулируемым порогом;",
            "зеркальное отражение по горизонтали;",
            "поворот изображения на 90 градусов по часовой стрелке.",
        ],
    )
    add_text_paragraph(
        document,
        "Программа должна поддерживать открытие распространенных растровых форматов, по меньшей мере PNG, JPEG, BMP, GIF и TIFF, "
        "а также сохранение результата в форматы PNG, JPEG и BMP."
    )

    add_heading(document, "Требования к надежности", level=2)
    add_bullets(
        document,
        [
            "приложение не должно аварийно завершаться при выборе неподдерживаемого или поврежденного файла;",
            "при ошибке открытия или сохранения пользователь должен получать понятное диагностическое сообщение;",
            "алгоритмы обработки должны работать с копией изображения, не изменяя исходные данные без явного подтверждения пользователя;",
            "ядро обработки изображений должно покрываться модульными тестами для ключевых операций.",
        ],
    )

    add_heading(document, "Условия эксплуатации", level=2)
    add_text_paragraph(
        document,
        "Эксплуатация программы предполагается на персональном компьютере или ноутбуке под управлением операционной системы Windows 10 "
        "или Windows 11 с установленной платформой .NET Desktop Runtime совместимой версии."
    )
    add_text_paragraph(
        document,
        "Пользователь должен обладать базовыми навыками работы с оконным интерфейсом, файловыми диалогами и каталогами хранения документов."
    )

    add_heading(document, "Требования к составу и параметрам технических средств", level=2)
    add_bullets(
        document,
        [
            "процессор архитектуры x64 с тактовой частотой не ниже 1,8 ГГц;",
            "оперативная память не менее 4 ГБ;",
            "не менее 200 МБ свободного места на диске для приложения и временных файлов;",
            "монитор с разрешением не ниже 1280 x 720 пикселей;",
            "устройства ввода: клавиатура и мышь или функционально аналогичное указательное устройство.",
        ],
    )

    add_heading(document, "Требования к информационной и программной совместимости", level=2)
    add_bullets(
        document,
        [
            "приложение должно быть реализовано на платформе .NET с использованием WPF для пользовательского интерфейса;",
            "алгоритмы обработки изображений должны быть выделены в отдельную библиотеку классов без привязки к WPF;",
            "проект тестирования должен использовать фреймворк xUnit;",
            "структура решения должна обеспечивать раздельную сборку основного приложения, библиотеки и тестового проекта.",
        ],
    )

    add_heading(document, "Требования к маркировке и упаковке", level=2)
    add_text_paragraph(
        document,
        "Специальные требования к маркировке и упаковке не предъявляются, поскольку программное средство распространяется в электронном виде."
    )

    add_heading(document, "Требования к транспортированию и хранению", level=2)
    add_text_paragraph(
        document,
        "Транспортирование и хранение осуществляются в форме электронного копирования файлов проекта и публикационных артефактов. "
        "Следует обеспечивать сохранность исходного кода, документации и сборок на надежных носителях или в системе контроля версий."
    )

    add_heading(document, "Специальные требования", level=2)
    add_bullets(
        document,
        [
            "пользовательский интерфейс должен быть русскоязычным;",
            "основные действия пользователя должны выполняться не более чем за 2-3 последовательных шага;",
            "решение должно допускать дальнейшее расширение списка операций обработки без полной переработки архитектуры.",
        ],
    )

    add_heading(document, "Требования к программной документации", level=1)
    add_text_paragraph(document, "В состав программной документации должны входить:")
    add_bullets(
        document,
        [
            "настоящее техническое задание;",
            "исходный код программного средства;",
            "пояснительная записка по проекту;",
            "материалы по тестированию и результаты модульных тестов;",
            "краткое описание пользовательского сценария работы с программой.",
        ],
    )

    add_heading(document, "Технико-экономические показатели", level=1)
    add_text_paragraph(
        document,
        "Разрабатываемое программное средство сокращает время выполнения базовых операций обработки изображений по сравнению с ручным "
        "подходом и обеспечивает единообразие результата при демонстрации алгоритмов обработки."
    )
    add_text_paragraph(
        document,
        "Использование выделенной библиотеки алгоритмов и автоматизированных тестов снижает трудоемкость сопровождения, упрощает "
        "расширение функциональности и повышает повторное использование кода."
    )

    add_heading(document, "Стадии и этапы разработки", level=1)
    add_stages_table(document)

    add_heading(document, "Порядок контроля и приемки", level=1)
    add_text_paragraph(
        document,
        "Контроль разработки и приемка программного средства выполняются по результатам анализа исходного кода, сборки решения, прохождения "
        "модульных тестов и ручной проверки ключевых сценариев работы."
    )
    add_text_paragraph(document, "Приемка считается успешной при выполнении следующих условий:")
    add_bullets(
        document,
        [
            "решение успешно собирается без ошибок;",
            "модульные тесты для ядра обработки изображений проходят успешно;",
            "приложение корректно открывает поддерживаемые изображения;",
            "каждая заявленная операция обработки формирует ожидаемый визуальный результат;",
            "сохранение обработанного изображения выполняется без потери работоспособности приложения;",
            "комплект документации сформирован и доступен для последующего оформления.",
        ],
    )
    add_text_paragraph(
        document,
        "При обнаружении дефектов приемка откладывается до устранения замечаний и повторного проведения контрольных испытаний."
    )

    add_heading(document, "Источники разработки", level=1)
    add_bullets(
        document,
        [
            "задание на тему «создать программное средство для обработки изображений»;",
            "ГОСТ 19.201-78 «Техническое задание. Требования к содержанию и оформлению»;",
            "ГОСТ 19.101-77 «Виды программ и программных документов»;",
            "документация платформы .NET и WPF для выбранного технологического стека.",
        ],
    )


def update_fields_and_export_pdf(docx_path: Path, pdf_path: Path) -> None:
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0

    try:
        document = word.Documents.Open(str(docx_path))
        document.Fields.Update()
        for toc in document.TablesOfContents:
            toc.Update()
        document.Save()
        document.ExportAsFixedFormat(str(pdf_path), 17)
        document.Close(False)
    finally:
        word.Quit()


def main() -> None:
    ensure_directories()
    create_base_document_from_template()

    document = Document(BASE_DOCX_PATH)
    clear_document_body(document)
    document.core_properties.title = "Техническое задание на разработку ImageToolkit"
    document.core_properties.subject = "Техническое задание"
    document.core_properties.author = "Codex"

    add_title_page(document)
    clone_last_section(document)
    add_main_content(document)

    document.save(OUTPUT_DOCX_PATH)
    update_fields_and_export_pdf(OUTPUT_DOCX_PATH, OUTPUT_PDF_PATH)

    print(OUTPUT_DOCX_PATH)
    print(OUTPUT_PDF_PATH)


if __name__ == "__main__":
    main()
